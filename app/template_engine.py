"""
Engine điền dữ liệu hóa đơn vào BẤT KỲ template Excel nào, dựa trên placeholder.

Cách viết template (mở bằng Excel, gõ thẳng vào ô):

  * Trường đơn:      {{so_hoa_don}}, {{ben_ban.ten_cong_ty}}, {{ngay.dd}} ...
    - Ô chỉ chứa đúng 1 placeholder  -> ghi giá trị gốc (số vẫn là số).
    - Ô trộn chữ + placeholder        -> thay thế thành chuỗi.
  * Dòng hàng mục:   một dòng có chứa {{item.xxx}} (vd {{item.ten_hang_hoa_dich_vu}}).
    Dòng này được nhân bản theo số hàng mục trên hóa đơn; mọi thứ bên dưới
    (công thức, ô gộp, chiều cao dòng, ảnh, vùng in) tự dịch xuống.
  * Trường tự nhập:  {{extra.ten_bat_ky}} - lấy từ tham số `extra` khi gọi API
    (người đề nghị, bộ phận, số đề nghị, người ký ...).
  * Sheet có tên bắt đầu bằng "_" (vd "_Huong dan") sẽ bị xóa khỏi file kết quả.

Công thức có sẵn trong template (vd =SUM(K16:K18)) được giữ và tự giãn
như khi chèn dòng trong Excel. Lưu ý: công thức chỉ hiện kết quả khi Excel tính lại
(KHÔNG hiện trong Protected View, khung xem trước của Outlook/Gmail, điện thoại).
Với số tiền, nên dùng placeholder ({{tong_cong}}, {{so_tien_bo_sung}}) để luôn hiện đúng.
"""
from __future__ import annotations

import io
import re
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.formula.tokenizer import Token, Tokenizer
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import CellRange

from .context import KNOWN_ITEM_KEYS, build_context, build_item_context, is_known_key

DEFAULT_TEMPLATE_PATH = Path(__file__).resolve().parent / "templates" / "default.xlsx"

PLACEHOLDER = re.compile(r"\{\{\s*([A-Za-z_][\w.]*)\s*\}\}")
FULL_PLACEHOLDER = re.compile(r"^\s*\{\{\s*([A-Za-z_][\w.]*)\s*\}\}\s*$")
_CELL_PART = re.compile(r"^(\$?[A-Za-z]{0,3})(\$?)(\d+)$")


class TemplateError(ValueError):
    """Template không hợp lệ (người dùng cần sửa template)."""


# ---------------------------------------------------------------------------
# Dịch tham chiếu dòng trong công thức khi chèn k dòng ngay sau dòng `row`
# ---------------------------------------------------------------------------
def _shift_part(part: str, row: int, k: int, is_end: bool, range_start: int | None) -> str:
    m = _CELL_PART.match(part)
    if not m:
        return part  # cột trơn (A:A), tên vùng ... -> giữ nguyên
    col, dollar, num = m.group(1), m.group(2), int(m.group(3))
    if num > row or (is_end and num == row and range_start is not None and range_start <= row):
        num += k
    return f"{col}{dollar}{num}"


def _shift_ref(ref: str, row: int, k: int, sheet_title: str) -> str:
    sheet_prefix = ""
    body = ref
    if "!" in ref:
        sheet_prefix, body = ref.rsplit("!", 1)
        name = sheet_prefix.strip("'").replace("''", "'")
        if name != sheet_title:
            return ref  # tham chiếu sang sheet khác -> không đụng
        sheet_prefix += "!"
    if ":" in body:
        start, end = body.split(":", 1)
        ms = _CELL_PART.match(start)
        start_row = int(ms.group(3)) if ms else None
        new = (
            _shift_part(start, row, k, False, None)
            + ":"
            + _shift_part(end, row, k, True, start_row)
        )
    else:
        new = _shift_part(body, row, k, False, None)
    return sheet_prefix + new


def shift_formula(formula: str, row: int, k: int, sheet_title: str) -> str:
    if not isinstance(formula, str) or not formula.startswith("="):
        return formula
    tok = Tokenizer(formula)
    out = []
    for t in tok.items:
        if t.type == Token.OPERAND and t.subtype == Token.RANGE:
            out.append(_shift_ref(t.value, row, k, sheet_title))
        else:
            out.append(t.value)
    return "=" + "".join(out)


# ---------------------------------------------------------------------------
# Chèn k dòng sau dòng `row`, dịch mọi thứ bên dưới
# ---------------------------------------------------------------------------
def _insert_rows_after(ws, row: int, k: int):
    if k <= 0:
        return
    max_row, max_col = ws.max_row, ws.max_column
    title = ws.title

    # 1. Công thức trong toàn bộ workbook trỏ tới sheet này
    for sheet in ws.parent.worksheets:
        for r in sheet.iter_rows():
            for c in r:
                if isinstance(c.value, str) and c.value.startswith("="):
                    if sheet is ws:
                        c.value = shift_formula(c.value, row, k, title)
                    elif title in c.value:
                        c.value = _shift_cross_sheet(c.value, row, k, title)

    # 2. Ô gộp: tháo ra, ghi nhớ, gộp lại sau khi dịch
    affected = []
    for mr in list(ws.merged_cells.ranges):
        if mr.max_row > row:
            affected.append(CellRange(mr.coord))
            ws.unmerge_cells(mr.coord)

    # 3. Dịch giá trị + style
    if max_row > row:
        ws.move_range(
            f"A{row + 1}:{get_column_letter(max_col)}{max_row}", rows=k, translate=False
        )

    for cr in affected:
        if cr.min_row > row:
            cr.shift(row_shift=k)
        else:  # vùng gộp chứa dòng hàng mục và kéo dài xuống dưới -> giãn ra
            cr.expand(down=k)
        ws.merge_cells(cr.coord)

    # 4. Chiều cao dòng
    dims = ws.row_dimensions
    for r in range(max_row, row, -1):
        src = dims[r]
        dst = dims[r + k]
        dst.height = src.height
        dst.hidden = src.hidden
    for r in range(row + 1, row + k + 1):
        dims[r].height = dims[row].height
        dims[r].hidden = dims[row].hidden

    # 5. Ảnh (logo, chữ ký scan ...)
    for img in getattr(ws, "_images", []):
        anchor = img.anchor
        for marker_name in ("_from", "to"):
            marker = getattr(anchor, marker_name, None)
            if marker is not None and marker.row >= row:  # marker.row tính từ 0
                marker.row += k

    # 6. Vùng in
    if ws.print_area:
        areas = []
        for area in ws.print_area.split(","):
            ref = area.split("!")[-1].replace("$", "")
            areas.append(_shift_ref(ref, row, k, title))
        ws.print_area = ",".join(areas)

    # 7. Data validation
    for dv in ws.data_validations.dataValidation:
        new_ranges = []
        for cr in dv.sqref.ranges:
            cr = CellRange(cr.coord)
            if cr.min_row > row:
                cr.shift(row_shift=k)
            elif cr.max_row >= row:
                cr.expand(down=k)
            new_ranges.append(cr.coord)
        dv.sqref = " ".join(new_ranges)


def _shift_cross_sheet(formula: str, row: int, k: int, title: str) -> str:
    tok = Tokenizer(formula)
    out = []
    for t in tok.items:
        if t.type == Token.OPERAND and t.subtype == Token.RANGE and "!" in t.value:
            out.append(_shift_ref(t.value, row, k, title))
        else:
            out.append(t.value)
    return "=" + "".join(out)


def _copy_row(ws, src_row: int, dst_row: int):
    for col in range(1, ws.max_column + 1):
        src = ws.cell(row=src_row, column=col)
        dst = ws.cell(row=dst_row, column=col)
        if src.has_style:
            dst._style = copy(src._style)
        v = src.value
        if isinstance(v, str) and v.startswith("="):
            v = Translator(v, origin=src.coordinate).translate_formula(dst.coordinate)
        dst.value = v
    for mr in list(ws.merged_cells.ranges):
        if mr.min_row == mr.max_row == src_row:
            ws.merge_cells(
                start_row=dst_row, start_column=mr.min_col,
                end_row=dst_row, end_column=mr.max_col,
            )


# ---------------------------------------------------------------------------
# Thay placeholder
# ---------------------------------------------------------------------------
def _to_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _render(value, lookup):
    """Trả về (giá_trị_mới, các_key_đã_dùng)."""
    if not isinstance(value, str) or "{{" not in value:
        return value, []
    m = FULL_PLACEHOLDER.match(value)
    if m:
        return lookup(m.group(1)), [m.group(1)]
    keys = []
    in_formula = value.lstrip().startswith("=")

    def sub(mm):
        keys.append(mm.group(1))
        text = _to_text(lookup(mm.group(1)))
        # placeholder đặt trong chuỗi của công thức mẫu (="Số "&"{{so_hoa_don}}"): nhân đôi dấu " để
        # dữ liệu không thoát ra khỏi chuỗi và chèn thêm hàm (HYPERLINK, WEBSERVICE...)
        return text.replace('"', '""') if in_formula else text

    text = PLACEHOLDER.sub(sub, value)
    return (text if text.strip() else None), keys


def _apply(cell, lookup):
    """Thay placeholder trong ô. Chống chèn công thức (formula injection): chỉ ô nào trong MẪU vốn đã là
    công thức mới được là công thức; giá trị lấy từ hóa đơn/người nhập bắt đầu bằng "=" luôn ghi là chữ."""
    original = cell.value
    cell.value, _ = _render(original, lookup)
    if isinstance(cell.value, str) and cell.value.startswith("=") and not original.lstrip().startswith("="):
        cell.data_type = "s"   # openpyxl tự coi chuỗi "=..." là công thức -> ép về kiểu chữ


def _find_item_row(ws) -> int | None:
    rows = set()
    for r in ws.iter_rows():
        for c in r:
            if isinstance(c.value, str) and "{{" in c.value:
                if any(k.startswith("item.") for k in PLACEHOLDER.findall(c.value)):
                    rows.add(c.row)
    if len(rows) > 1:
        raise TemplateError(
            f"Sheet '{ws.title}' có placeholder {{{{item.*}}}} ở nhiều dòng "
            f"({sorted(rows)}). Chỉ được đặt trên MỘT dòng hàng mục."
        )
    return rows.pop() if rows else None


def _item_columns(ws, row: int) -> dict[str, int]:
    """Map key item -> cột, cho các ô chỉ chứa đúng 1 placeholder item."""
    cols = {}
    for c in ws[row]:
        if isinstance(c.value, str):
            m = FULL_PLACEHOLDER.match(c.value)
            if m and m.group(1).startswith("item."):
                cols[m.group(1)[5:]] = c.column
    return cols


def _fill_sheet(ws, ctx: dict, items: list[dict], warnings: list[str]):
    unknown = set()

    def make_lookup(item_ctx=None):
        def lookup(key):
            if key.startswith("item."):
                sub = key[5:]
                if item_ctx is None or sub not in KNOWN_ITEM_KEYS:
                    unknown.add(key)
                    return None
                return item_ctx.get(sub)
            if key.startswith("extra."):
                return ctx["extra"].get(key[6:])
            if not is_known_key(key):
                unknown.add(key)
                return None
            value = ctx["flat"].get(key)
            if callable(value):  # giá trị sinh lười (vd số đề nghị) - chỉ tính 1 lần
                value = ctx["flat"][key] = value()
            return value
        return lookup

    item_row = _find_item_row(ws)
    if item_row is not None:
        n = max(len(items), 1)
        _insert_rows_after(ws, item_row, n - 1)
        for i in range(1, n):
            _copy_row(ws, item_row, item_row + i)

        for i in range(n):
            r = item_row + i
            item = items[i] if i < len(items) else {}
            ictx = build_item_context(item, i) if items else {}
            lookup = make_lookup(ictx)
            for c in ws[r]:
                if isinstance(c.value, str) and "{{" in c.value:
                    _apply(c, lookup)

    lookup = make_lookup(None)
    for r in ws.iter_rows():
        for c in r:
            if isinstance(c.value, str) and "{{" in c.value:
                _apply(c, lookup)

    for key in sorted(unknown):
        warnings.append(f"Template dùng placeholder không tồn tại: {{{{{key}}}}} (sheet '{ws.title}')")


# ---------------------------------------------------------------------------
# Thay logo
# ---------------------------------------------------------------------------
_EMU_PER_PT = 12700


def _row_height_emu(ws, row0: int) -> int:
    """Chiều cao dòng (row0 tính từ 0) theo EMU; dòng chưa đặt chiều cao = 15pt."""
    h = ws.row_dimensions[row0 + 1].height
    return int((h if h is not None else 15) * _EMU_PER_PT)


def _anchor_height_emu(ws, anchor) -> int | None:
    """Chiều cao đang hiển thị của ảnh neo kiểu twoCell / oneCell."""
    ext = getattr(anchor, "ext", None)
    if ext is not None and getattr(ext, "height", None):
        return int(ext.height)
    fr, to = getattr(anchor, "_from", None), getattr(anchor, "to", None)
    if fr is None or to is None:
        return None
    total = sum(_row_height_emu(ws, r) for r in range(fr.row, to.row))
    return total - fr.rowOff + to.rowOff


def replace_logo(wb, logo: bytes, warnings: list[str]) -> bool:
    """Thay ảnh nằm trên cùng-bên trái của sheet đầu tiên bằng `logo`.
    Giữ nguyên vị trí và chiều cao ảnh cũ; chiều rộng theo đúng tỉ lệ logo mới (không méo)."""
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
    from openpyxl.drawing.xdr import XDRPositiveSize2D

    ws = wb.worksheets[0]
    images = [im for im in getattr(ws, "_images", []) if getattr(im.anchor, "_from", None) is not None]
    if not images:
        warnings.append("Mẫu Excel không có ảnh logo nên không thay được logo đã chọn.")
        return False
    old = min(images, key=lambda im: (im.anchor._from.row, im.anchor._from.col))

    new = XLImage(io.BytesIO(logo))
    height = _anchor_height_emu(ws, old.anchor) or int(new.height * 9525)  # 1px = 9525 EMU
    width = int(height * new.width / new.height)
    fr = old.anchor._from
    new.anchor = OneCellAnchor(
        _from=AnchorMarker(col=fr.col, colOff=fr.colOff, row=fr.row, rowOff=fr.rowOff),
        ext=XDRPositiveSize2D(width, height),
    )
    ws._images[ws._images.index(old)] = new
    return True


def load_template(template_bytes: bytes | None = None):
    try:
        if template_bytes:
            return openpyxl.load_workbook(io.BytesIO(template_bytes))
        return openpyxl.load_workbook(DEFAULT_TEMPLATE_PATH)
    except Exception as e:  # noqa: BLE001
        raise TemplateError(f"Không mở được file template .xlsx: {e}") from e


def fill_template(
    data: dict,
    template_bytes: bytes | None = None,
    extra: dict | None = None,
    so_de_nghi_provider=None,
    logo: bytes | None = None,
) -> tuple[bytes, list[str]]:
    """Điền `data` vào template. Trả về (nội dung file .xlsx, danh sách cảnh báo).
    `logo`: ảnh PNG/JPG thay cho logo trên cùng của sheet đầu tiên (None = giữ logo của mẫu)."""
    wb = load_template(template_bytes)
    ctx, items, warnings = build_context(data, extra or {}, so_de_nghi_provider)
    if logo:
        replace_logo(wb, logo, warnings)

    for title in list(wb.sheetnames):
        if title.startswith("_"):
            del wb[title]
    if not wb.sheetnames:
        raise TemplateError("Template không còn sheet nào sau khi bỏ các sheet '_...'.")

    if template_bytes is None or b"xl/media/" in template_bytes:
        try:
            import PIL  # noqa: F401  openpyxl cần Pillow để giữ logo/ảnh trong mẫu
        except ImportError:
            warnings.append("Máy chủ thiếu thư viện Pillow nên logo/ảnh trong mẫu Excel bị mất. "
                            "Chạy: pip install pillow rồi khởi động lại server.")

    found_any = False
    for ws in wb.worksheets:
        if any(isinstance(c.value, str) and "{{" in c.value for r in ws.iter_rows() for c in r):
            found_any = True
            _fill_sheet(ws, ctx, items, warnings)
    if not found_any:
        warnings.append("Template không chứa placeholder {{...}} nào - file trả về giữ nguyên.")

    # Công thức còn lại trong mẫu (vd =TODAY()) được Excel tính lại ngay khi mở file
    wb.calculation.fullCalcOnLoad = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), warnings
