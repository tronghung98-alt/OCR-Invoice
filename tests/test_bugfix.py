"""Test hồi quy cho danh sách bug đợt kiểm thử 26/10/2026 (Danh_sach_bug.xlsx)."""
import io
import os
import subprocess
import sys

import openpyxl
import pytest

from app import config, security, services
from app.checks import check_invoice
from app.context import to_number, to_number_vn
from app.db import SessionLocal
from app.models import REVIEW, PaymentRequest, fold
from app.template_engine import fill_template
from app.web import parse_tax_rate, parse_tax_rates
from tests.test_engine import SAMPLE, custom_template
from tests.test_web import client, ai, csrf_of, login, make_user, one, run_ai, upload  # noqa: F401


def edit_form(client, rid, **over):
    """Form trang Sửa với 1 dòng hàng; `over` ghi đè từng ô."""
    url = f"/de-nghi/{rid}/sua"
    form = {"csrf": csrf_of(client, url), "so_hoa_don": "123", "ngay_thang_nam": "05/10/2026",
            "ben_ban_ten_cong_ty": "Công ty ABC", "thue_suat": "8%", "tien_thue": "",
            "item_ten": ["Hàng"], "item_sl": ["1"], "item_dg": ["100.000"], "item_tt": [""],
            "item_dvt": [""], "item_gc": [""], "then": "stay"}
    form.update(over)
    return client.post(url, data=form, follow_redirects=False)


def reviewed_request(client):
    make_user("an")
    login(client, "an")
    upload(client)
    run_ai()
    r = one()
    assert r.status == REVIEW
    return r


# ---------------------------------------------------------------- BUG-01
@pytest.mark.parametrize("raw, want", [("0,125", 0.125), ("0.125", 0.125), ("-0,5", -0.5), (".5", 0.5),
                                       ("2.300", 2300), ("2.300.000", 2300000), ("1,5", 1.5)])
def test_bug01_to_number_small_decimal(raw, want):
    assert to_number(raw) == want


@pytest.mark.parametrize("raw, want", [("0,125", 0.125), ("1,5", 1.5), ("12,345", 12.345), ("1.500", 1500),
                                       ("1.250,5", 1250.5), ("3", 3), (2.5, 2.5)])
def test_bug01_quantity_uses_vietnamese_decimal_comma(raw, want):
    assert to_number_vn(raw) == want


def test_bug01_decimal_quantity_end_to_end(client):
    r = reviewed_request(client)
    assert edit_form(client, r.id, item_sl=["0,125"], item_dg=["80.000"], item_tt=[""]).status_code == 303
    item = one().data["hang_muc"][0]
    assert (item["so_luong"], item["thanh_tien"]) == (0.125, 10000)
    ws = openpyxl.load_workbook(io.BytesIO(client.get(f"/de-nghi/{r.id}/excel").content)).active
    assert ws["I16"].value == 0.125 and ws["K16"].value == 10000


def test_bug01_large_quantity_is_flagged():
    data = dict(SAMPLE, hang_muc=[{"ten_hang_hoa_dich_vu": "Ốc", "so_luong": 125, "don_gia": 80000,
                                   "thanh_tien": 10000000}], tong_tien_truoc_thue=10000000, tien_thue=800000,
                tong_tien_thanh_toan=10800000)
    assert not any(i["code"] == "qty_large" for i in check_invoice(data))
    data["hang_muc"][0].update(so_luong=1250, thanh_tien=100000000)
    data.update(tong_tien_truoc_thue=100000000, tien_thue=8000000, tong_tien_thanh_toan=108000000)
    issue = next(i for i in check_invoice(data) if i["code"] == "qty_large")
    assert issue["level"] == "info" and "1.250" in issue["message"]


# ---------------------------------------------------------------- BUG-02
def test_bug02_suite_is_isolated_from_dotenv():
    assert config.DATABASE_URL == os.environ["DATABASE_URL"] and not config.SMTP_HOST
    assert not os.getenv("ADMIN_USERNAME") and not os.getenv("GEMINI_API_KEY")


def test_bug02_env_vars_win_over_dotenv(tmp_path):
    """.env không còn ghi đè biến môi trường đã đặt; DOTENV_DISABLED=1 bỏ qua hẳn .env."""
    env_file = tmp_path / ".env"
    # chạy config với BASE_DIR trỏ tới thư mục tạm có .env giả
    script = f"""
import os, sys, pathlib, runpy
pathlib.Path(r'{env_file}').write_text('SMTP_HOST=smtp.that.vn\\nX_ONLY=1\\n', encoding='utf-8')
import dotenv
real = dotenv.load_dotenv
dotenv.load_dotenv = lambda path, override=False: real(r'{env_file}', override=override)
import app.config as c
print(repr(c.SMTP_HOST), repr(os.getenv('X_ONLY')))
"""
    base = {k: v for k, v in os.environ.items() if k not in ("DOTENV_DISABLED", "SMTP_HOST", "X_ONLY")}
    run = lambda env: subprocess.run([sys.executable, "-c", script], env=env, capture_output=True,  # noqa: E731
                                     text=True, cwd=config.BASE_DIR, check=True).stdout.split()
    assert run({**base, "SMTP_HOST": ""}) == ["''", "'1'"]               # biến đã đặt (kể cả rỗng) thắng .env
    assert run({**base, "DOTENV_DISABLED": "1"}) == ["''", "None"]       # bỏ qua hẳn .env
    assert run(base) == ["'smtp.that.vn'", "'1'"]                         # không đặt gì -> lấy từ .env


# ---------------------------------------------------------------- BUG-03
@pytest.mark.parametrize("evil", ['=HYPERLINK("http://evil.example","Bam vao")', "=1+1", "=cmd|' /C calc'!A0"])
def test_bug03_formula_injection_written_as_text(evil):
    data = dict(SAMPLE, ben_ban={"ten_cong_ty": evil},
                hang_muc=[dict(SAMPLE["hang_muc"][0], ten_hang_hoa_dich_vu=evil)])
    content, _ = fill_template(data, None, {"nguoi_de_nghi": evil}, lambda: "X1")
    ws = openpyxl.load_workbook(io.BytesIO(content)).active
    cells = [c for row in ws.iter_rows() for c in row if isinstance(c.value, str) and "evil" in c.value or
             (isinstance(c.value, str) and c.value in ("=1+1", "=cmd|' /C calc'!A0"))]
    assert cells, "phải tìm thấy giá trị đã điền"
    for c in cells:
        assert c.data_type == "s", f"{c.coordinate} bị ghi thành công thức"
    assert ws["C8"].data_type == "f"  # công thức có sẵn trong mẫu vẫn là công thức


def test_bug03_placeholder_inside_template_formula_cannot_break_out():
    tpl = custom_template(A1='="Số HĐ: "&"{{so_hoa_don}}"')
    content, _ = fill_template(dict(SAMPLE, so_hoa_don='1"&HYPERLINK("http://evil")&"'), tpl, {})
    v = openpyxl.load_workbook(io.BytesIO(content)).active["A1"].value
    assert v == '="Số HĐ: "&"1""&HYPERLINK(""http://evil"")&"""'


# ---------------------------------------------------------------- BUG-04
def test_bug04_tax_rate_parser():
    assert parse_tax_rates("5%, 10%") == [0.05, 0.1] and parse_tax_rates("5% và 10%") == [0.05, 0.1]
    assert parse_tax_rates("KCT, 8%") == [0.0, 0.08] and parse_tax_rates("10,5%") == [0.105]
    assert parse_tax_rates("abc") is None and parse_tax_rates("10% abc") is None
    assert parse_tax_rate("10%") == 0.1 and parse_tax_rate("KCT") == 0.0 and parse_tax_rate("5%, 10%") is None


def test_bug04_mixed_tax_rate_invoice_can_be_saved(client):
    r = reviewed_request(client)
    two_rows = {"item_ten": ["Sách", "Dịch vụ"], "item_sl": ["1", "1"], "item_dg": ["100.000", "200.000"],
                "item_tt": ["", ""], "item_dvt": ["", ""], "item_gc": ["", ""]}
    # nhiều mức thuế + tiền thuế theo hóa đơn (5% x 100.000 + 10% x 200.000 = 25.000) -> lưu được, giữ đúng số
    assert edit_form(client, r.id, thue_suat="5%, 10%", tien_thue="25.000", **two_rows).status_code == 303
    d = one().data
    assert (d["tong_tien_truoc_thue"], d["tien_thue"], d["tong_tien_thanh_toan"]) == (300000, 25000, 325000)
    assert not any("thuế" in w.lower() for w in one().warnings)
    # nhiều mức nhưng bỏ trống tiền thuế -> báo lỗi rõ ràng
    page = edit_form(client, r.id, thue_suat="5%, 10%", tien_thue="", **two_rows)
    assert page.status_code == 400 and "nhiều mức thuế" in page.text
    # một mức: server vẫn tự tính, bỏ qua số trình duyệt gửi
    assert edit_form(client, r.id, thue_suat="10%", tien_thue="1", **two_rows).status_code == 303
    assert one().data["tien_thue"] == 30000


def test_bug04_multi_rate_check_uses_range():
    base = dict(SAMPLE, thue_suat="5%, 10%")
    assert not any(i["code"] == "tax_rate" for i in check_invoice(dict(base, tien_thue=20000,
                                                                          tong_tien_thanh_toan=280000)))
    out = [i for i in check_invoice(dict(base, tien_thue=90000, tong_tien_thanh_toan=350000)) if i["code"] == "tax_rate"]
    assert out and "khoảng" in out[0]["message"]


# ---------------------------------------------------------------- BUG-05
def test_bug05_fold():
    assert fold("Công ty ĐIỆN LỰC Hà Nội") == "cong ty dien luc ha noi"


@pytest.mark.parametrize("q", ["Điện lực", "điện lực", "ĐIỆN LỰC", "dien luc", "HÀ NỘI"])
def test_bug05_search_vietnamese_case_insensitive(client, ai, q):
    ai.results = [dict(SAMPLE, ben_ban={"ten_cong_ty": "Công ty Điện lực Hà Nội"})]
    make_user("an")
    login(client, "an")
    upload(client)
    run_ai()
    r = one()
    assert r.id and f"/de-nghi/{r.id}" in client.get("/de-nghi", params={"q": q}).text
    assert f"/de-nghi/{r.id}\"" not in client.get("/de-nghi", params={"q": "nuoc sach"}).text


def test_bug05_search_wildcards_are_literal(client):
    r = reviewed_request(client)
    assert f"/de-nghi/{r.id}" not in client.get("/de-nghi", params={"q": "%"}).text


def test_bug05_startup_backfills_old_rows(client):
    r = reviewed_request(client)
    with SessionLocal() as db:
        db.execute(PaymentRequest.__table__.update().values(search_text=None))
        db.commit()
    services.on_startup()
    with SessionLocal() as db:
        assert "cong ty abc" in db.get(PaymentRequest, r.id).search_text


# ---------------------------------------------------------------- BUG-07
def test_bug07_one_wrong_attempt_after_lock_expires_does_not_relock(monkeypatch):
    security._fails.clear()
    now = [1000.0]
    monkeypatch.setattr(security.time, "time", lambda: now[0])
    for _ in range(security.MAX_FAILS):
        security.record_login("1.1.1.1", "an", False)
    assert security.login_locked("1.1.1.1", "an") > 0
    now[0] += security.LOCK_SECONDS + 1
    assert security.login_locked("1.1.1.1", "an") == 0
    security.record_login("1.1.1.1", "an", False)          # sai 1 lần sau khi hết khóa
    assert security.login_locked("1.1.1.1", "an") == 0      # chưa bị khóa lại
    for _ in range(security.MAX_FAILS - 1):
        security.record_login("1.1.1.1", "an", False)
    assert security.login_locked("1.1.1.1", "an") > 0      # đủ 5 lần mới khóa


# ---------------------------------------------------------------- BUG-08
@pytest.mark.parametrize("name, ok", [("nguyen.van-a_1", True), ("nguyễn", False), ("an an", False),
                                      ("", False), ("a" * 65, False), ("đức", False)])
def test_bug08_username_rule(name, ok):
    assert (security.validate_username(name) is None) is ok


def test_bug08_create_user_rejects_vietnamese_letters(client):
    make_user("boss", admin=True)
    login(client, "boss")
    client.post("/nguoi-dung", data={"csrf": csrf_of(client, "/nguoi-dung"), "username": "nguyễn",
                                     "password": "matkhau123"})
    with SessionLocal() as db:
        from app.models import User
        assert db.query(User).filter(User.username == "nguyễn").count() == 0


def test_bug08_cli_rejects_vietnamese_letters(monkeypatch):
    from app import manage
    monkeypatch.setattr(manage, "_ask_password", lambda: "matkhau123")
    with pytest.raises(SystemExit) as e:
        manage.main(["create-user", "nguyễn"])
    assert "không dấu" in str(e.value)


# ---------------------------------------------------------------- BUG-09
def test_bug09_no_file_selected_friendly_error(client):
    make_user("an")
    login(client, "an")
    r = client.post("/de-nghi", data={"csrf": csrf_of(client, "/")})
    assert r.status_code == 400 and "Chọn ít nhất một hóa đơn." in r.text
