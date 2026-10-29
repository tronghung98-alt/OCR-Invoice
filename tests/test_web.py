"""Test toàn bộ quy trình theo sơ đồ: đăng nhập -> gửi nhiều hóa đơn -> vòng lặp AI
(đọc, tự kiểm tra, đọc lại, xuất Excel) -> xem và sửa -> gửi mail -> quản lý xác nhận."""
import email
import io
import re
from datetime import timedelta
from email import policy

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import text

from app import config, services
from app.db import SessionLocal, engine
from app.extractor import ExtractionError
from app.main import app
from app.models import (CONFIRMED, NEEDS_HUMAN, QUEUED, REVIEW, SENT, PaymentRequest, User, utcnow)
from app.security import _fails, hash_password
from app.worker import worker
from tests.test_engine import SAMPLE

PDF = b"%PDF-1.7 fake"


class FakeAI:
    """Giả lập Gemini: trả lần lượt các kết quả trong `results` (dict hoặc Exception)."""

    def __init__(self, *results):
        self.results, self.calls = list(results), []

    def __call__(self, content, mime, hint=None):
        self.calls.append(hint)
        r = self.results.pop(0) if len(self.results) > 1 else self.results[0]
        if isinstance(r, Exception):
            raise r
        return dict(r)


@pytest.fixture
def ai(monkeypatch):
    fake = FakeAI(SAMPLE)
    monkeypatch.setattr(services, "extract_invoice_data", fake)
    return fake


@pytest.fixture
def client(ai, tmp_path, monkeypatch):
    _fails.clear()
    monkeypatch.setattr(config, "OUTBOX_DIR", tmp_path / "outbox")
    with TestClient(app) as c:
        yield c


def run_ai():
    while worker.run_once():
        pass


def make_user(username, admin=False, manager=False, email_addr=None, pw="matkhau123"):
    with SessionLocal() as db:
        u = User(username=username, password_hash=hash_password(pw), full_name=username.title(),
                 department="Kanri", is_admin=admin, is_manager=manager, email=email_addr)
        db.add(u)
        db.commit()
        return u.id


def csrf_of(c, path="/login"):
    return re.search(r'name="csrf" value="([^"]+)"', c.get(path).text).group(1)


def login(c, username, pw="matkhau123"):
    return c.post("/login", data={"username": username, "password": pw, "csrf": csrf_of(c), "next": "/de-nghi"},
                  follow_redirects=False)


def upload(c, files=(("hd.pdf", PDF),), **fields):
    return c.post("/de-nghi", data={"csrf": csrf_of(c, "/"), **fields},
                  files=[("files", f) for f in files], follow_redirects=False)


def one():
    with SessionLocal() as db:
        return db.query(PaymentRequest).order_by(PaymentRequest.id.desc()).first()


def event_kinds(rid):
    with SessionLocal() as db:
        return [e.kind for e in db.get(PaymentRequest, rid).events]


def set_field(**values):
    with SessionLocal() as db:
        for r in db.query(PaymentRequest):
            for k, v in values.items():
                setattr(r, k, v)
        db.commit()


def sent_mails():
    return sorted(config.OUTBOX_DIR.glob("*.eml")) if config.OUTBOX_DIR.exists() else []


# ---------------------------------------------------------------- đăng nhập
def test_login_shows_sent_list(client):
    make_user("an")
    r = login(client, "an")
    assert r.status_code == 303 and r.headers["location"] == "/de-nghi"
    assert "Đề nghị đã gửi" in client.get("/de-nghi").text


def test_pages_require_login_and_csrf(client):
    for path in ("/", "/de-nghi", "/de-nghi/1", "/de-nghi/1/sua", "/de-nghi/1/gui", "/nguoi-dung"):
        r = client.get(path, follow_redirects=False)
        assert r.status_code == 303 and r.headers["location"].startswith("/login?next=")
    make_user("an")
    login(client, "an")
    assert client.post("/de-nghi", data={}, files=[("files", ("a.pdf", PDF))]).status_code == 400
    for _ in range(5):
        login(client, "an", "sai")
    assert login(client, "an").status_code == 429


# ---------------------------------------------------------------- vòng lặp AI
def test_upload_many_then_ai_loop_reads_each(client, ai):
    make_user("an")
    login(client, "an")
    r = upload(client, files=(("a.pdf", PDF), ("b.png", b"\x89PNG\r\n\x1a\nxx")), nguoi_de_nghi="An")
    assert r.headers["location"] == "/de-nghi?tab=ai-dang-doc"
    with SessionLocal() as db:
        assert [x.status for x in db.query(PaymentRequest)] == [QUEUED, QUEUED]
    # màn hình chờ khi AI đang đọc phiếu của chính người dùng (BUG-10: thay cho data-autorefresh cũ)
    page = client.get("/de-nghi").text
    assert 'id="busy-overlay"' in page and 'data-poll-url="/trang-thai-xu-ly"' in page
    assert "Còn 2 phiếu đang xử lý" in page

    run_ai()  # 1. nhận -> 2. đọc -> 4. tự kiểm tra -> 3. xuất Excel -> 5. báo cáo
    with SessionLocal() as db:
        reqs = db.query(PaymentRequest).order_by(PaymentRequest.id).all()
        assert [x.status for x in reqs] == [REVIEW, REVIEW]
        assert [x.so_de_nghi[:2] for x in reqs] == ["01", "02"]
        assert all(x.ai_reads == 1 and x.warnings == [] for x in reqs)
        kinds = [e.kind for e in reqs[0].events]
    assert kinds == ["upload", "ai_read", "review"]
    page = client.get("/de-nghi?tab=can-xu-ly").text
    assert page.count('data-href="/de-nghi/') == 2
    assert not worker.run_once()  # hết hóa đơn -> chờ


def test_self_check_rereads_with_hint_and_keeps_best(client, monkeypatch):
    wrong = dict(SAMPLE, tong_tien_thanh_toan=2808000)  # AI đọc nhầm dấu phân cách
    fake = FakeAI(wrong, SAMPLE)
    monkeypatch.setattr(services, "extract_invoice_data", fake)
    make_user("an")
    login(client, "an")
    upload(client)
    run_ai()
    r = one()
    assert r.status == REVIEW and r.ai_reads == 2 and r.warnings == []
    assert fake.calls[0] is None and "khác tổng thanh toán" in fake.calls[1][0]
    assert r.tong_tien == 280800


def test_still_suspect_after_reread_goes_to_review_with_warnings(client, monkeypatch):
    monkeypatch.setattr(services, "extract_invoice_data", FakeAI(dict(SAMPLE, tong_tien_thanh_toan=999)))
    make_user("an")
    login(client, "an")
    upload(client)
    run_ai()
    r = one()
    assert r.status == REVIEW and r.ai_reads == 2
    assert any("khác tổng thanh toán" in w for w in r.warnings)
    assert "điểm cần xem lại" in client.get(f"/de-nghi/{r.id}").text


def test_blurry_image_needs_human_then_human_types_data(client, monkeypatch):
    fake = FakeAI({"kha_nang_doc": "khong_doc_duoc", "hang_muc": []})
    monkeypatch.setattr(services, "extract_invoice_data", fake)
    make_user("an")
    login(client, "an")
    upload(client)
    run_ai()
    r = one()
    assert r.status == NEEDS_HUMAN and r.ai_reads == 1 and "mờ" in r.error
    assert r.output_path is None
    page = client.get(f"/de-nghi/{r.id}")  # dữ liệu thiếu hầu hết các trường vẫn phải hiển thị được
    assert page.status_code == 200 and "Cần người xử lý" in page.text and "Nhập hoặc sửa tay" in page.text
    assert client.get(f"/de-nghi/{r.id}/sua").status_code == 200
    assert client.get(f"/de-nghi/{r.id}/gui", follow_redirects=False).status_code == 303  # chưa gửi được

    url = f"/de-nghi/{r.id}/sua"
    form = {"csrf": csrf_of(client, url), "so_hoa_don": "777", "ngay_thang_nam": "1/10/2026",
            "ben_ban_ten_cong_ty": "Cty XYZ", "thue_suat": "8%", "tien_thue": "8.000",
            "tong_tien_truoc_thue": "100.000", "tong_tien_thanh_toan": "108.000",
            "item_ten": ["Bút", ""], "item_sl": ["10", ""], "item_dg": ["10.000", ""], "item_tt": ["100.000", ""],
            "item_dvt": ["cây", ""], "item_gc": ["", ""], "then": "stay"}
    assert client.post(url, data=form, follow_redirects=False).status_code == 303
    r = one()
    assert r.status == REVIEW and r.edited_by_human and r.warnings == []
    assert r.data["ngay_thang_nam"] == "01/10/2026" and r.data["hang_muc"][0]["don_gia"] == 10000
    ws = openpyxl.load_workbook(io.BytesIO(client.get(f"/de-nghi/{r.id}/excel").content)).active
    assert ws["D16"].value == "Bút" and ws["K16"].value == 100000  # số tĩnh, không công thức (BUG-06)


def test_edit_validation_errors_keep_data(client):
    make_user("an")
    login(client, "an")
    upload(client)
    run_ai()
    r = one()
    url = f"/de-nghi/{r.id}/sua"
    res = client.post(url, data={"csrf": csrf_of(client, url), "item_ten": ["X"], "item_tt": ["abc"],
                                 "ngay_thang_nam": "45/13/2026", "tong_tien_thanh_toan": "nhiều"})
    assert res.status_code == 400
    assert "phải là số" in res.text and "Ngày hóa đơn không hợp lệ" in res.text
    assert one().data == r.data  # không lưu gì


def test_transient_error_retries_next_round_then_needs_human(client, monkeypatch):
    monkeypatch.setattr(services, "extract_invoice_data", FakeAI(ExtractionError("503 UNAVAILABLE", transient=True)))
    monkeypatch.setattr(config, "MAX_TRANSIENT_ATTEMPTS", 3)
    make_user("an")
    login(client, "an")
    upload(client)
    run_ai()
    r = one()
    assert r.status == QUEUED and r.attempts == 1 and "thử lại" in r.error
    assert not worker.run_once()  # chưa tới giờ thử lại -> AI chờ
    for _ in range(2):
        set_field(next_attempt_at=utcnow() - timedelta(seconds=1))
        run_ai()
    r = one()
    assert r.status == NEEDS_HUMAN and "liên tục" in r.error


def test_permanent_error_needs_human(client, monkeypatch):
    monkeypatch.setattr(services, "extract_invoice_data", FakeAI(ExtractionError("404 model not found")))
    make_user("an")
    login(client, "an")
    upload(client)
    run_ai()
    assert one().status == NEEDS_HUMAN


def test_reread_button_B(client, monkeypatch):
    make_user("an")
    login(client, "an")
    upload(client)
    run_ai()
    r = one()
    first_no = r.so_de_nghi
    monkeypatch.setattr(services, "extract_invoice_data", FakeAI(dict(SAMPLE, so_hoa_don="999")))
    client.post(f"/de-nghi/{r.id}/doc-lai", data={"csrf": csrf_of(client, f"/de-nghi/{r.id}")})
    assert one().status == QUEUED
    run_ai()
    r = one()
    assert r.status == REVIEW and r.so_hoa_don == "999" and r.so_de_nghi == first_no  # giữ số đề nghị


# ---------------------------------------------------------------- gửi mail & quản lý xác nhận
def _send(client, rid, to="sep@cty.vn", **extra):
    url = f"/de-nghi/{rid}/gui"
    return client.post(url, data={"csrf": csrf_of(client, url), "to": to, "note": "Nhờ anh duyệt", **extra},
                       follow_redirects=False)


def _token_from_last_mail():
    msg = email.message_from_bytes(sent_mails()[-1].read_bytes(), policy=policy.default)
    body = msg.get_body(("plain",)).get_content()
    return msg, re.search(r"http://dntt\.local/xac-nhan/([\w-]+)", body).group(1)


def test_full_flow_send_and_manager_confirms(client):
    make_user("an")
    make_user("sep", manager=True, email_addr="sep@cty.vn")
    login(client, "an")
    upload(client)
    run_ai()
    r = one()
    page = client.get(f"/de-nghi/{r.id}/gui").text
    assert "sep@cty.vn" in page and "quay lại sửa" in page  # A

    assert _send(client, r.id).status_code == 303
    r = one()
    assert r.status == SENT and r.sent_to == "sep@cty.vn"
    msg, token = _token_from_last_mail()
    assert msg["To"] == "sep@cty.vn" and r.so_de_nghi in msg["Subject"]
    att = next(msg.iter_attachments())
    assert att.get_filename().endswith(".xlsx") and att.get_content()[:2] == b"PK"

    anon = TestClient(app)  # quản lý không đăng nhập
    page = anon.get(f"/xac-nhan/{token}")
    assert page.status_code == 200 and "Xác nhận đã đọc" in page.text
    assert one().status == SENT  # chỉ mở link (GET) thì chưa xác nhận
    assert anon.get(f"/xac-nhan/{token}/excel").content[:2] == b"PK"
    tok = re.search(r'name="csrf" value="([^"]+)"', page.text).group(1)
    anon.post(f"/xac-nhan/{token}", data={"csrf": tok})
    r = one()
    assert r.status == CONFIRMED and r.confirmed_by == "sep@cty.vn"
    assert "Đã xác nhận đọc" in anon.get(f"/xac-nhan/{token}").text
    lst = client.get("/de-nghi?tab=da-xac-nhan").text
    assert r.so_de_nghi in lst
    assert event_kinds(r.id)[-2:] == ["send", "confirm"]
    # đã xác nhận -> khóa: không sửa, không đọc lại, không xóa (người thường)
    assert client.get(f"/de-nghi/{r.id}/sua", follow_redirects=False).status_code == 303
    client.post(f"/de-nghi/{r.id}/xoa", data={"csrf": csrf_of(client, f"/de-nghi/{r.id}")})
    assert one() is not None


def test_warnings_must_be_acknowledged_before_send(client, monkeypatch):
    monkeypatch.setattr(services, "extract_invoice_data", FakeAI(dict(SAMPLE, tong_tien_thanh_toan=999)))
    make_user("an")
    login(client, "an")
    upload(client)
    run_ai()
    rid = one().id
    _send(client, rid)
    assert one().status == REVIEW and not sent_mails()
    _send(client, rid, da_kiem_tra="true")
    assert one().status == SENT


def test_recall_and_resend_invalidate_old_link(client):
    make_user("an")
    login(client, "an")
    upload(client)
    run_ai()
    rid = one().id
    _send(client, rid)
    _, tok1 = _token_from_last_mail()
    _send(client, rid)  # gửi lại -> link cũ hết hiệu lực
    _, tok2 = _token_from_last_mail()
    anon = TestClient(app)
    assert anon.get(f"/xac-nhan/{tok1}").status_code == 404
    assert anon.get(f"/xac-nhan/{tok2}").status_code == 200

    r = client.post(f"/de-nghi/{rid}/thu-hoi", data={"csrf": csrf_of(client, f"/de-nghi/{rid}")},
                    follow_redirects=False)
    assert r.headers["location"] == f"/de-nghi/{rid}/sua"  # A: quay lại sửa
    assert one().status == REVIEW
    assert anon.get(f"/xac-nhan/{tok2}").status_code == 404


def test_expired_link_cannot_confirm(client):
    make_user("an")
    login(client, "an")
    upload(client)
    run_ai()
    _send(client, one().id)
    _, token = _token_from_last_mail()
    set_field(confirm_expires_at=utcnow() - timedelta(days=1))
    anon = TestClient(app)
    page = anon.get(f"/xac-nhan/{token}")
    assert "đã hết hạn" in page.text and "Xác nhận đã đọc</button>" not in page.text
    tok = csrf_of(anon, "/login")
    assert anon.post(f"/xac-nhan/{token}", data={"csrf": tok}).status_code == 410
    assert one().status == SENT


def test_bad_email_rejected(client):
    make_user("an")
    login(client, "an")
    upload(client)
    run_ai()
    _send(client, one().id, to="khong-phai-mail")
    assert one().status == REVIEW


# ---------------------------------------------------------------- phân quyền & quản trị
def test_permissions(client):
    make_user("an")
    make_user("binh")
    make_user("sep", admin=True)
    login(client, "an")
    upload(client)
    run_ai()
    rid = one().id
    client.cookies.clear()
    login(client, "binh")
    for path in (f"/de-nghi/{rid}", f"/de-nghi/{rid}/excel", f"/de-nghi/{rid}/sua", f"/de-nghi/{rid}/gui"):
        assert client.get(path).status_code == 404
    assert client.get("/nguoi-dung").status_code == 403
    client.cookies.clear()
    login(client, "sep")
    assert client.get(f"/de-nghi/{rid}").status_code == 200


def test_admin_sets_manager(client):
    make_user("sep", admin=True)
    uid = make_user("ql")
    login(client, "sep")
    tok = csrf_of(client, "/nguoi-dung")
    client.post(f"/nguoi-dung/{uid}", data={"action": "toggle_manager", "csrf": tok})
    with SessionLocal() as db:
        assert not db.get(User, uid).is_manager  # chưa có mail -> không cho
    client.post(f"/nguoi-dung/{uid}", data={"action": "set_email", "email": "ql@cty.vn", "csrf": tok})
    client.post(f"/nguoi-dung/{uid}", data={"action": "toggle_manager", "csrf": tok})
    with SessionLocal() as db:
        assert db.get(User, uid).is_manager


# ---------------------------------------------------------------- nâng cấp CSDL cũ
def test_upgrade_old_database_keeps_data():
    from app.migrate import upgrade
    from app.db import Base
    Base.metadata.drop_all(engine)
    with engine.begin() as c:  # cấu trúc bản trước (không có cột workflow, không có request_events)
        c.execute(text("CREATE TABLE users (id INTEGER PRIMARY KEY, username VARCHAR(64), password_hash VARCHAR(255),"
                       " full_name VARCHAR(128), department VARCHAR(128), is_admin BOOLEAN, is_active BOOLEAN,"
                       " created_at DATETIME, last_login_at DATETIME)"))
        c.execute(text("CREATE TABLE payment_requests (id INTEGER PRIMARY KEY, user_id INTEGER, status VARCHAR(16),"
                       " error TEXT, so_de_nghi VARCHAR(64), so_hoa_don VARCHAR(64), ngay_hoa_don VARCHAR(32),"
                       " ben_ban VARCHAR(255), tong_tien FLOAT, original_filename VARCHAR(255), mime_type VARCHAR(64),"
                       " original_path VARCHAR(512), template_path VARCHAR(512), output_path VARCHAR(512),"
                       " data JSON, extra JSON, warnings JSON, created_at DATETIME, updated_at DATETIME)"))
        c.execute(text("INSERT INTO users (id, username, password_hash, is_admin, is_active) VALUES (1,'cu','x',1,1)"))
        c.execute(text("INSERT INTO payment_requests (id, user_id, status, original_filename, mime_type, original_path)"
                       " VALUES (1,1,'done','a.pdf','application/pdf','files/1/a.pdf'),"
                       " (2,1,'failed','b.pdf','application/pdf','files/2/b.pdf')"))
    upgrade()
    upgrade()  # chạy lại không lỗi
    with SessionLocal() as db:
        a, b = db.get(PaymentRequest, 1), db.get(PaymentRequest, 2)
        assert (a.status, b.status) == (REVIEW, NEEDS_HUMAN)
        assert a.attempts == 0 and a.edited_by_human is False and list(a.events) == []
        u = db.get(User, 1)
        assert u.is_manager is False and u.email is None


def test_every_page_renders_in_every_status(client, monkeypatch):
    """Mở mọi trang ở mọi trạng thái - bắt lỗi template với dữ liệu thiếu."""
    make_user("an", admin=True)
    login(client, "an")
    monkeypatch.setattr(services, "extract_invoice_data", FakeAI({"kha_nang_doc": "khong_doc_duoc"}))
    upload(client)
    run_ai()  # needs_human
    monkeypatch.setattr(services, "extract_invoice_data", FakeAI({"hang_muc": [{"thanh_tien": "abc"}]}))
    upload(client)
    run_ai()  # dữ liệu rác
    monkeypatch.setattr(services, "extract_invoice_data", FakeAI(SAMPLE))
    upload(client)
    run_ai()
    _send(client, one().id)  # sent
    upload(client)  # queued
    with SessionLocal() as db:
        ids = [r.id for r in db.query(PaymentRequest)]
    for rid in ids:
        for path in ("", "/sua", "/gui", "/trang-thai"):
            res = client.get(f"/de-nghi/{rid}{path}")
            assert res.status_code == 200, (rid, path, res.status_code)
    for tab in ("", "can-xu-ly", "ai-dang-doc", "cho-quan-ly", "da-xac-nhan"):
        assert client.get(f"/de-nghi?tab={tab}").status_code == 200
    for path in ("/", "/tai-khoan", "/nguoi-dung"):
        assert client.get(path).status_code == 200


# ---------------------------------------------------------------- chọn logo
def _excel_logo_size(rid):
    with SessionLocal() as db:
        path = services.abs_path(db.get(PaymentRequest, rid).output_path)
    from PIL import Image
    ws = openpyxl.load_workbook(path).worksheets[0]
    return Image.open(io.BytesIO(ws._images[0]._data())).size


def test_choose_logo_on_upload_and_change_on_edit(client):
    from PIL import Image

    from app import logos
    size = {k: Image.open(io.BytesIO(logos.logo_bytes(k))).size for k in logos.LOGOS}
    make_user("an")
    login(client, "an")
    page = client.get("/").text
    assert 'name="logo" value="vietnam"' in page and 'name="logo" value="axlab"' in page

    upload(client, logo="vietnam")
    run_ai()
    r = one()
    assert r.status == REVIEW and r.extra["logo"] == "vietnam"
    assert _excel_logo_size(r.id) == size["vietnam"]

    # trang sửa hiện đúng logo đang chọn; đổi sang AX LAB rồi lưu
    edit = client.get(f"/de-nghi/{r.id}/sua").text
    assert re.search(r'value="vietnam"\s+checked', edit)
    form = {"csrf": csrf_of(client, f"/de-nghi/{r.id}/sua"), "logo": "axlab", "thue_suat": "8%",
            "ngay_thang_nam": "05/10/2026", "item_ten": ["RAM"], "item_sl": ["2"], "item_dg": ["100000"],
            "item_tt": ["200000"], "item_dvt": [""], "item_gc": [""], "then": "stay"}
    assert client.post(f"/de-nghi/{r.id}/sua", data=form, follow_redirects=False).status_code == 303
    assert one().extra["logo"] == "axlab" and _excel_logo_size(r.id) == size["axlab"]

    # giá trị lạ -> logo mặc định
    upload(client, logo="khong-co")
    run_ai()
    assert one().extra["logo"] == logos.DEFAULT_LOGO
